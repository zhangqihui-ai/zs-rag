"""Tests for MinerU routing in document_parser."""

from unittest.mock import MagicMock, patch

from app.core.document_parser import ParsedDocument, parse_document
from app.core.mineru_gateway import MineruResult
from app.core.parser_config import ParserOptions


def test_docx_engine_mineru_routes_to_gateway():
    fake_result = MineruResult(
        content_list=[{"type": "text", "text": "Hello", "page_idx": 0}],
        markdown="# Hello",
        source_file_name="sample.docx",
    )
    mock_gw = MagicMock()
    mock_gw.is_enabled.return_value = True
    mock_gw.should_handle.return_value = True
    mock_gw.backend = "pipeline"
    mock_gw.lang = "ch"
    mock_gw.parse.return_value = fake_result

    opts = ParserOptions(
        pdf_engine="opendataloader",
        pdf_hybrid=False,
        docx_engine="mineru",
        excel_engine="html_table",
        csv_engine="standard",
        text_engine="native",
    )

    with patch("app.core.mineru_gateway.get_mineru_gateway", return_value=mock_gw):
        doc = parse_document(
            "sample.docx",
            b"fake-docx-bytes",
            parser_options=opts,
        )

    assert doc.parser_type == "docx"
    assert doc.metadata.get("parser_backend") == "mineru"
    assert doc.metadata.get("_parse_markdown") == "# Hello"
    assert len(doc.metadata.get("_parse_content_list") or []) == 1
    mock_gw.parse.assert_called_once()


def test_docx_engine_mineru_skips_format_whitelist():
    fake_result = MineruResult(
        content_list=[{"type": "text", "text": "Hi", "page_idx": 0}],
        markdown="# Hi",
        source_file_name="a.docx",
    )
    mock_gw = MagicMock()
    mock_gw.is_enabled.return_value = True
    mock_gw.should_handle.return_value = False
    mock_gw.backend = "pipeline"
    mock_gw.lang = "ch"
    mock_gw.parse.return_value = fake_result

    opts = ParserOptions(
        pdf_engine="opendataloader",
        pdf_hybrid=False,
        docx_engine="mineru",
        excel_engine="html_table",
        csv_engine="standard",
        text_engine="native",
    )

    with patch("app.core.mineru_gateway.get_mineru_gateway", return_value=mock_gw):
        doc = parse_document("a.docx", b"bytes", parser_options=opts)

    assert doc.metadata.get("parser_backend") == "mineru"
    mock_gw.parse.assert_called_once()


def test_xls_excel_engine_mineru_falls_back_to_local_parser():
    mock_gw = MagicMock()
    mock_gw.is_enabled.return_value = True

    opts = ParserOptions(
        pdf_engine="opendataloader",
        pdf_hybrid=False,
        docx_engine="python-docx",
        excel_engine="mineru",
        csv_engine="standard",
        text_engine="native",
    )

    fake_doc = ParsedDocument(
        parser_type="xls",
        text="sheet data",
        char_count=10,
        segments=[],
        metadata={"excel_engine": "html_table"},
    )

    with patch("app.core.mineru_gateway.get_mineru_gateway", return_value=mock_gw):
        with patch("app.core.document_parser._parse_excel", return_value=fake_doc) as mock_excel:
            doc = parse_document("事项信息.xls", b"fake-xls", parser_options=opts)

    mock_gw.parse.assert_not_called()
    mock_excel.assert_called_once_with(b"fake-xls", "xls", engine="html_table", log=None)
    assert doc.metadata.get("parser_backend") == "html_table"
    assert doc.metadata.get("parser_fallback") is True
    assert doc.metadata.get("parser_primary") == "mineru"
    assert doc.parser_type == "xls"


def test_xlsx_backslash_zip_skips_mineru_for_local_parser():
    import zipfile
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "SheetA"
    ws["A1"] = "hello"
    good = BytesIO()
    wb.save(good)
    good_bytes = good.getvalue()

    bad = BytesIO()
    with zipfile.ZipFile(BytesIO(good_bytes), "r") as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item.filename.replace("/", "\\"), zin.read(item.filename))
    bad_bytes = bad.getvalue()

    mock_gw = MagicMock()
    mock_gw.is_enabled.return_value = True

    opts = ParserOptions(
        pdf_engine="opendataloader",
        pdf_hybrid=False,
        docx_engine="python-docx",
        excel_engine="mineru",
        csv_engine="standard",
        text_engine="native",
    )

    with patch("app.core.mineru_gateway.get_mineru_gateway", return_value=mock_gw):
        doc = parse_document("清单.xlsx", bad_bytes, parser_options=opts)

    mock_gw.parse.assert_not_called()
    assert doc.parser_type == "xlsx"
    assert "hello" in doc.text
